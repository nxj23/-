#!/usr/bin/env python3
"""dump 用户上传PPT的全部文本，定位需要回填实测数据的位置"""
from pptx import Presentation
from pptx.util import Emu

pptx_path = "/workspace/.uploads/dd0a6cb1-2836-4a1a-b573-0daa32a57d65_全链路性能瓶颈排查与综合调优_优化版.pptx"
prs = Presentation(pptx_path)
print(f"幻灯片总数: {len(prs.slides)}\n")

for i, slide in enumerate(prs.slides, 1):
    print(f"========== Slide {i} ==========")
    for shp in slide.shapes:
        if shp.has_text_frame:
            txt = shp.text_frame.text.strip()
            if txt:
                # 标注位置便于定位
                try:
                    x = round(shp.left/914400, 2) if shp.left else 0
                    y = round(shp.top/914400, 2) if shp.top else 0
                except:
                    x = y = 0
                print(f"  [{x:.1f},{y:.1f}] {repr(txt[:200])}")
        if shp.has_table:
            tbl = shp.table
            print(f"  -- 表格 {len(tbl.rows)}行x{len(tbl.columns)}列 --")
            for r in tbl.rows:
                cells = [c.text.strip().replace('\n',' ') for c in r.cells]
                print("    | " + " | ".join(cells))
    print()
